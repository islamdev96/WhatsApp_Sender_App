"""Contact file reading (CSV, Excel, TXT) and template generation."""
import csv
import os
import re
from utils.logger import logger, log_exception
from utils.helpers.phone import _normalize_phone


def read_contacts(file_path, default_country_code="20"):
    """Reads contacts from a CSV file and returns a list of dictionaries."""
    contacts = []
    try:
        if not os.path.exists(file_path):
            logger.warning(f"Contacts file not found: {file_path}")
            return []

        with open(file_path, 'r', encoding='utf-8-sig') as f:
            reader = csv.DictReader(f)
            for row in reader:
                normalized_row = {k.strip().lower(): v for k, v in row.items()}

                phone = normalized_row.get('phone') or normalized_row.get('mobile') or \
                        normalized_row.get('number') or normalized_row.get('phone 1 - value')
                name = normalized_row.get('name') or normalized_row.get('given name') or 'Customer'
                var1 = normalized_row.get('var1') or normalized_row.get('variable1') or normalized_row.get('v1')
                var2 = normalized_row.get('var2') or normalized_row.get('variable2') or normalized_row.get('v2')
                var3 = normalized_row.get('var3') or normalized_row.get('variable3') or normalized_row.get('v3')
                var4 = normalized_row.get('var4') or normalized_row.get('variable4') or normalized_row.get('v4')
                var5 = normalized_row.get('var5') or normalized_row.get('variable5') or normalized_row.get('v5')

                phone = _normalize_phone(phone, default_country_code)
                if phone:
                    c = {'phone': phone, 'name': name.strip()}
                    if var1: c['var1'] = var1
                    if var2: c['var2'] = var2
                    if var3: c['var3'] = var3
                    if var4: c['var4'] = var4
                    if var5: c['var5'] = var5
                    contacts.append(c)
        logger.info(f"Loaded {len(contacts)} contacts from {file_path}")
    except FileNotFoundError:
        logger.error(f"Contacts file not found: {file_path}")
    except csv.Error as exc:
        logger.error(f"CSV parsing error in {file_path}: {exc}")
    except Exception as exc:
        log_exception(f"Error reading contacts from {file_path}", exc)

    return contacts



def read_contacts_excel(file_path, default_country_code="20"):
    """Reads contacts from an Excel (.xlsx) file."""
    contacts = []
    try:
        import openpyxl
        if not os.path.exists(file_path):
            logger.warning(f"Excel file not found: {file_path}")
            return []

        wb = openpyxl.load_workbook(file_path, read_only=True)
        ws = wb.active

        # Read header row
        headers = []
        for cell in next(ws.iter_rows(min_row=1, max_row=1)):
            headers.append(str(cell.value or '').strip().lower())

        # Find phone and name columns
        phone_col = None
        name_col = None
        for i, h in enumerate(headers):
            if h in ('phone', 'mobile', 'number', 'phone 1 - value', 'رقم', 'هاتف'):
                phone_col = i
            if h in ('name', 'given name', 'اسم', 'الاسم'):
                name_col = i
        var_cols = {}
        for i, h in enumerate(headers):
            if h in ('var1', 'variable1', 'v1'):
                var_cols['var1'] = i
            if h in ('var2', 'variable2', 'v2'):
                var_cols['var2'] = i
            if h in ('var3', 'variable3', 'v3'):
                var_cols['var3'] = i
            if h in ('var4', 'variable4', 'v4'):
                var_cols['var4'] = i
            if h in ('var5', 'variable5', 'v5'):
                var_cols['var5'] = i

        if phone_col is None:
            # Try first two columns: assume col 0=name, col 1=phone
            if len(headers) >= 2:
                name_col = 0
                phone_col = 1
            elif len(headers) == 1:
                phone_col = 0

        if phone_col is None:
            logger.warning(f"No phone column found in Excel file: {file_path}")
            return []

        for row in ws.iter_rows(min_row=2):
            phone_val = row[phone_col].value if phone_col < len(row) else None
            name_val = row[name_col].value if name_col is not None and name_col < len(row) else 'Customer'

            phone = _normalize_phone(phone_val, default_country_code)
            if phone:
                c = {'phone': phone, 'name': str(name_val or 'Customer').strip()}
                for k, idx in var_cols.items():
                    if idx < len(row):
                        val = row[idx].value
                        if val is not None:
                            c[k] = str(val)
                contacts.append(c)

        wb.close()
        logger.info(f"Loaded {len(contacts)} contacts from Excel file: {file_path}")
    except FileNotFoundError:
        logger.error(f"Excel file not found: {file_path}")
    except ImportError:
        logger.error("openpyxl library is not installed, cannot read Excel files")
    except Exception as exc:
        log_exception(f"Error reading Excel file {file_path}", exc)

    return contacts



def read_contacts_txt(file_path, default_country_code="20"):
    """Reads contacts from a plain text file, extracting phone numbers.
    
    Supports two formats:
    1. One number per line (most common)
    2. Numbers mixed with text (extracted via regex)
    """
    contacts = []
    try:
        if not os.path.exists(file_path):
            logger.warning(f"Text file not found: {file_path}")
            return []
        
        # Try UTF-8 first, then fallback to cp1256 (Arabic Windows)
        content = None
        for encoding in ('utf-8', 'utf-8-sig', 'cp1256', 'latin-1'):
            try:
                with open(file_path, 'r', encoding=encoding) as f:
                    content = f.read()
                break
            except (UnicodeDecodeError, UnicodeError):
                continue
        
        if not content:
            logger.warning(f"Could not read text file {file_path} with any supported encoding")
            return []
        
        seen = set()
        
        # Strategy 1: Try line-by-line (most common for phone lists)
        for line in content.splitlines():
            line = line.strip()
            if not line:
                continue
            # If the line is purely digits (with optional +, spaces, dashes, dots)
            cleaned = re.sub(r'[\s\-\.\(\)\+]+', '', line)
            if cleaned.isdigit() and 8 <= len(cleaned) <= 15:
                phone = _normalize_phone(line, default_country_code)
                if phone and phone not in seen:
                    seen.add(phone)
                    contacts.append({'phone': phone, 'name': 'عميل'})
        
        # Strategy 2: If no numbers found line-by-line, try regex extraction
        if not contacts:
            matches = re.findall(r'\+?\d(?:[\d\-\s\.]*\d){8,14}', content)
            for match in matches:
                phone = _normalize_phone(match, default_country_code)
                if phone and phone not in seen:
                    seen.add(phone)
                    contacts.append({'phone': phone, 'name': 'عميل'})

        logger.info(f"Loaded {len(contacts)} contacts from text file: {file_path}")
    except FileNotFoundError:
        logger.error(f"Text file not found: {file_path}")
    except Exception as exc:
        log_exception(f"Error reading text file {file_path}", exc)

    return contacts



def read_contacts_auto(file_path, default_country_code="20"):
    """Auto-detect file type and read contacts accordingly."""
    if not file_path or not os.path.exists(file_path):
        return []
    ext = os.path.splitext(file_path)[1].lower()
    if ext in ('.xlsx', '.xls'):
        return read_contacts_excel(file_path, default_country_code)
    elif ext == '.txt':
        return read_contacts_txt(file_path, default_country_code)
    else:
        return read_contacts(file_path, default_country_code)



def create_contacts_template(file_path):
    """Creates a sample CSV template for the user."""
    try:
        with open(file_path, 'w', newline='', encoding='utf-8-sig') as f:
            writer = csv.writer(f)
            writer.writerow(['Name', 'Phone', 'Var1', 'Var2', 'Var3', 'Var4', 'Var5'])
            writer.writerow(['Client Name', '010XXXXXXXX', 'Value1', 'Value2', 'Value3', 'Value4', 'Value5'])
        return True
    except OSError as exc:
        logger.error(f"Could not create contacts template {file_path}: {exc}")
        return False
    except Exception as exc:
        log_exception(f"Unexpected error creating contacts template {file_path}", exc)
        return False



