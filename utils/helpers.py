import csv
import os
import re
from utils.logger import logger, log_exception


def _normalize_phone(phone, default_country_code="20"):
    """Normalize a phone number string.

    Handles:
    - Leading '+' sign (strips it for WhatsApp API)
    - Numbers already starting with country code
    - Egypt-specific local formats (01X → 201X) when default is '20'
    - Generic local numbers with configurable default country code
    """
    if not phone:
        return None
    phone = str(phone).strip()
    # Remove .0 from float conversion
    if phone.endswith('.0'):
        phone = phone[:-2]
    # Remove common formatting chars
    phone = re.sub(r'[\s\-\(\)\.]+', '', phone)
    # Handle + prefix
    if phone.startswith('+'):
        phone = phone[1:]
    
    # Smart generic normalizing: if a local number starts with 0 (e.g. 05xxx or 01xxx)
    # and default_country_code is set, strip the leading 0 and prepend the country code.
    if default_country_code:
        default_cc = str(default_country_code).strip().replace("+", "")
        if phone.startswith('0') and not phone.startswith('00') and len(phone) > 4:
            phone = default_cc + phone[1:]

    # If already starts with country code, return as-is
    if default_country_code and phone.startswith(str(default_country_code)):
        return phone
    # Egypt-specific: local mobile numbers
    if default_country_code == "20":
        if phone.startswith('01') and len(phone) == 11:
            return '2' + phone
        if phone.startswith('1') and len(phone) == 10:
            return '20' + phone
    # Generic: prepend default country code for short local numbers
    if default_country_code and len(phone) <= 10 and not phone.startswith('0'):
        return str(default_country_code) + phone
    return phone


def normalize_phone(phone, default_country_code="20"):
    """Public wrapper for phone normalization."""
    return _normalize_phone(phone, default_country_code)


def read_contacts(file_path, default_country_code="20"):
    """Reads contacts from a CSV file and returns a list of dictionaries."""
    contacts = []
    try:
        if not os.path.exists(file_path):
            logger.warning(f"Contacts file not found: {file_path}")
            return []

        with open(file_path, 'r', encoding='utf-8-sig') as f:
            reader = csv.DictReader(f)
            for row in reader:
                normalized_row = {k.strip().lower(): v for k, v in row.items()}

                phone = normalized_row.get('phone') or normalized_row.get('mobile') or \
                        normalized_row.get('number') or normalized_row.get('phone 1 - value')
                name = normalized_row.get('name') or normalized_row.get('given name') or 'Customer'
                var1 = normalized_row.get('var1') or normalized_row.get('variable1') or normalized_row.get('v1')
                var2 = normalized_row.get('var2') or normalized_row.get('variable2') or normalized_row.get('v2')
                var3 = normalized_row.get('var3') or normalized_row.get('variable3') or normalized_row.get('v3')
                var4 = normalized_row.get('var4') or normalized_row.get('variable4') or normalized_row.get('v4')
                var5 = normalized_row.get('var5') or normalized_row.get('variable5') or normalized_row.get('v5')

                phone = _normalize_phone(phone, default_country_code)
                if phone:
                    c = {'phone': phone, 'name': name.strip()}
                    if var1: c['var1'] = var1
                    if var2: c['var2'] = var2
                    if var3: c['var3'] = var3
                    if var4: c['var4'] = var4
                    if var5: c['var5'] = var5
                    contacts.append(c)
        logger.info(f"Loaded {len(contacts)} contacts from {file_path}")
    except FileNotFoundError as exc:
        logger.error(f"Contacts file not found: {file_path}")
    except csv.Error as exc:
        logger.error(f"CSV parsing error in {file_path}: {exc}")
    except Exception as exc:
        log_exception(f"Error reading contacts from {file_path}", exc)

    return contacts


def read_contacts_excel(file_path, default_country_code="20"):
    """Reads contacts from an Excel (.xlsx) file."""
    contacts = []
    try:
        import openpyxl
        if not os.path.exists(file_path):
            logger.warning(f"Excel file not found: {file_path}")
            return []

        wb = openpyxl.load_workbook(file_path, read_only=True)
        ws = wb.active

        # Read header row
        headers = []
        for cell in next(ws.iter_rows(min_row=1, max_row=1)):
            headers.append(str(cell.value or '').strip().lower())

        # Find phone and name columns
        phone_col = None
        name_col = None
        for i, h in enumerate(headers):
            if h in ('phone', 'mobile', 'number', 'phone 1 - value', 'رقم', 'هاتف'):
                phone_col = i
            if h in ('name', 'given name', 'اسم', 'الاسم'):
                name_col = i
        var_cols = {}
        for i, h in enumerate(headers):
            if h in ('var1', 'variable1', 'v1'):
                var_cols['var1'] = i
            if h in ('var2', 'variable2', 'v2'):
                var_cols['var2'] = i
            if h in ('var3', 'variable3', 'v3'):
                var_cols['var3'] = i
            if h in ('var4', 'variable4', 'v4'):
                var_cols['var4'] = i
            if h in ('var5', 'variable5', 'v5'):
                var_cols['var5'] = i

        if phone_col is None:
            # Try first two columns: assume col 0=name, col 1=phone
            if len(headers) >= 2:
                name_col = 0
                phone_col = 1
            elif len(headers) == 1:
                phone_col = 0

        if phone_col is None:
            logger.warning(f"No phone column found in Excel file: {file_path}")
            return []

        for row in ws.iter_rows(min_row=2):
            phone_val = row[phone_col].value if phone_col < len(row) else None
            name_val = row[name_col].value if name_col is not None and name_col < len(row) else 'Customer'

            phone = _normalize_phone(phone_val, default_country_code)
            if phone:
                c = {'phone': phone, 'name': str(name_val or 'Customer').strip()}
                for k, idx in var_cols.items():
                    if idx < len(row):
                        val = row[idx].value
                        if val is not None:
                            c[k] = str(val)
                contacts.append(c)

        wb.close()
        logger.info(f"Loaded {len(contacts)} contacts from Excel file: {file_path}")
    except FileNotFoundError as exc:
        logger.error(f"Excel file not found: {file_path}")
    except ImportError:
        logger.error("openpyxl library is not installed, cannot read Excel files")
    except Exception as exc:
        log_exception(f"Error reading Excel file {file_path}", exc)

    return contacts


def read_contacts_txt(file_path, default_country_code="20"):
    """Reads contacts from a plain text file, extracting phone numbers.
    
    Supports two formats:
    1. One number per line (most common)
    2. Numbers mixed with text (extracted via regex)
    """
    contacts = []
    try:
        if not os.path.exists(file_path):
            logger.warning(f"Text file not found: {file_path}")
            return []
        
        # Try UTF-8 first, then fallback to cp1256 (Arabic Windows)
        content = None
        for encoding in ('utf-8', 'utf-8-sig', 'cp1256', 'latin-1'):
            try:
                with open(file_path, 'r', encoding=encoding) as f:
                    content = f.read()
                break
            except (UnicodeDecodeError, UnicodeError):
                continue
        
        if not content:
            logger.warning(f"Could not read text file {file_path} with any supported encoding")
            return []
        
        seen = set()
        
        # Strategy 1: Try line-by-line (most common for phone lists)
        for line in content.splitlines():
            line = line.strip()
            if not line:
                continue
            # If the line is purely digits (with optional +, spaces, dashes, dots)
            cleaned = re.sub(r'[\s\-\.\(\)\+]+', '', line)
            if cleaned.isdigit() and 8 <= len(cleaned) <= 15:
                phone = _normalize_phone(line, default_country_code)
                if phone and phone not in seen:
                    seen.add(phone)
                    contacts.append({'phone': phone, 'name': 'عميل'})
        
        # Strategy 2: If no numbers found line-by-line, try regex extraction
        if not contacts:
            matches = re.findall(r'\+?\d(?:[\d\-\s\.]*\d){8,14}', content)
            for match in matches:
                phone = _normalize_phone(match, default_country_code)
                if phone and phone not in seen:
                    seen.add(phone)
                    contacts.append({'phone': phone, 'name': 'عميل'})

        logger.info(f"Loaded {len(contacts)} contacts from text file: {file_path}")
    except FileNotFoundError as exc:
        logger.error(f"Text file not found: {file_path}")
    except Exception as exc:
        log_exception(f"Error reading text file {file_path}", exc)

    return contacts


def read_contacts_auto(file_path, default_country_code="20"):
    """Auto-detect file type and read contacts accordingly."""
    if not file_path or not os.path.exists(file_path):
        return []
    ext = os.path.splitext(file_path)[1].lower()
    if ext in ('.xlsx', '.xls'):
        return read_contacts_excel(file_path, default_country_code)
    elif ext == '.txt':
        return read_contacts_txt(file_path, default_country_code)
    else:
        return read_contacts(file_path, default_country_code)


def create_contacts_template(file_path):
    """Creates a sample CSV template for the user."""
    try:
        with open(file_path, 'w', newline='', encoding='utf-8-sig') as f:
            writer = csv.writer(f)
            writer.writerow(['Name', 'Phone', 'Var1', 'Var2', 'Var3', 'Var4', 'Var5'])
            writer.writerow(['Client Name', '010XXXXXXXX', 'Value1', 'Value2', 'Value3', 'Value4', 'Value5'])
        return True
    except OSError as exc:
        logger.error(f"Could not create contacts template {file_path}: {exc}")
        return False
    except Exception as exc:
        log_exception(f"Unexpected error creating contacts template {file_path}", exc)
        return False


def check_proxy(proxy_type, host, port, username=None, password=None, timeout=10):
    """Tests a proxy connection using urllib.
    
    Returns (success, info_dict).
    """
    import urllib.request
    import urllib.error
    import json

    proxy_type = proxy_type.lower()
    proxy_url = f"{proxy_type}://"
    if username and password:
        proxy_url += f"{username}:{password}@"
    proxy_url += f"{host}:{port}"
    
    proxy_handler = urllib.request.ProxyHandler({
        'http': proxy_url,
        'https': proxy_url
    })
    
    opener = urllib.request.build_opener(proxy_handler)
    try:
        # Use http://ip-api.com/json (clean HTTP) to verify public IP and location details
        response = opener.open("http://ip-api.com/json", timeout=timeout)
        data = json.loads(response.read().decode('utf-8'))
        if data.get('status') == 'success':
            return True, {
                'ip': data.get('query'),
                'country': data.get('country'),
                'city': data.get('city'),
                'isp': data.get('isp')
            }
        else:
            return True, {
                'ip': data.get('query') or 'Unknown',
                'country': 'Unknown',
                'city': 'Unknown',
                'isp': 'Unknown'
            }
    except Exception as e:
        return False, {'error': str(e)}


def create_proxy_extension(profile_dir, proxy_type, host, port, username, password):
    """Generates a custom Chrome extension dynamically to handle proxy credentials authentication."""
    import json
    ext_dir = os.path.join(profile_dir, "proxy_extension")
    os.makedirs(ext_dir, exist_ok=True)
    
    manifest_path = os.path.join(ext_dir, "manifest.json")
    background_path = os.path.join(ext_dir, "background.js")
    
    manifest_json = {
        "version": "1.0.0",
        "manifest_version": 2,
        "name": "Chrome Proxy Helper Extension",
        "permissions": [
            "proxy",
            "tabs",
            "unlimitedStorage",
            "storage",
            "<all_urls>",
            "webRequest",
            "webRequestBlocking"
        ],
        "background": {
            "scripts": ["background.js"]
        },
        "minimum_chrome_version": "22.0.0"
    }
    
    background_js = f"""
    var config = {{
        mode: "fixed_servers",
        rules: {{
            singleProxy: {{
                scheme: "{proxy_type.lower()}",
                host: "{host}",
                port: parseInt({port})
            }},
            bypassList: []
        }}
    }};

    chrome.proxy.settings.set({{value: config, scope: "regular"}}, function() {{}});

    chrome.webRequest.onAuthRequired.addListener(
        function(details) {{
            return {{
                authCredentials: {{
                    username: "{username}",
                    password: "{password}"
                }}
            }};
        }},
        {{urls: ["<all_urls>"]}},
        ["blocking"]
    );
    """
    
    with open(manifest_path, "w", encoding="utf-8") as f:
        json.dump(manifest_json, f, indent=4)
        
    with open(background_path, "w", encoding="utf-8") as f:
        f.write(background_js)
        
    return ext_dir


def generate_random_fingerprint():
    """Generates a random desktop browser user-agent and resolution footprint."""
    import random
    user_agents = [
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36",
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/123.0.0.0 Safari/537.36",
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36",
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64; rv:124.0) Gecko/20100101 Firefox/124.0",
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64; rv:123.0) Gecko/20100101 Firefox/123.0",
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36 Edg/124.0.0.0",
        "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36",
        "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/605.1.15 (KHTML, like Gecko) Version/17.4 Safari/605.1.15"
    ]
    resolutions = [
        "1920,1080",
        "1366,768",
        "1440,900",
        "1536,864",
        "1600,900"
    ]
    return {
        "user_agent": random.choice(user_agents),
        "resolution": random.choice(resolutions)
    }


def cleanup_proxy_extension(profile_dir):
    """Remove the dynamically generated proxy extension from a Chrome profile.

    Should be called after the browser session ends to avoid leaving
    stale credential files on disk.
    """
    import shutil
    ext_dir = os.path.join(profile_dir, "proxy_extension")
    if os.path.isdir(ext_dir):
        try:
            shutil.rmtree(ext_dir, ignore_errors=True)
            logger.info("Cleaned up proxy extension in %s", profile_dir)
        except Exception as exc:
            log_exception(f"Could not remove proxy extension dir {ext_dir}", exc)


def cleanup_old_reports(reports_base_dir=None, max_age_days=30):
    """Delete report files (CSV/TXT) older than *max_age_days*.

    Scans `reports/` and `reports/number_checks/` for stale files
    and removes them to prevent unbounded disk usage.

    Returns the number of files removed.
    """
    import time

    if reports_base_dir is None:
        reports_base_dir = os.path.join(os.getcwd(), "reports")
    if not os.path.isdir(reports_base_dir):
        return 0

    cutoff = time.time() - (max_age_days * 86400)
    removed = 0
    scan_dirs = [reports_base_dir]

    # Also scan known subdirectories
    for sub in ("number_checks", "logs"):
        sub_path = os.path.join(reports_base_dir, sub)
        if os.path.isdir(sub_path):
            scan_dirs.append(sub_path)

    for dir_path in scan_dirs:
        try:
            for entry in os.scandir(dir_path):
                if not entry.is_file():
                    continue
                ext = os.path.splitext(entry.name)[1].lower()
                if ext not in (".csv", ".txt", ".log"):
                    continue
                try:
                    if entry.stat().st_mtime < cutoff:
                        os.unlink(entry.path)
                        removed += 1
                except OSError as exc:
                    logger.debug("Could not remove old report %s: %s", entry.path, exc)
        except OSError as exc:
            logger.debug("Could not scan reports directory %s: %s", dir_path, exc)

    if removed:
        logger.info("Cleaned up %d old report file(s) from %s", removed, reports_base_dir)
    return removed
